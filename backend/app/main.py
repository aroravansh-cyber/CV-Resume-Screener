# ============================================================
# ResumeAI - FastAPI Backend
# COMPLETE REPLACEMENT main.py
# ============================================================

from fastapi import (
    FastAPI,
    UploadFile,
    File,
    HTTPException,
    BackgroundTasks,
    Request,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
from pathlib import Path
from datetime import datetime
from typing import Optional, List, Dict, Any

import uuid
import threading
import re
import subprocess
import traceback


# ============================================================
# OPTIONAL LIBRARIES
# ============================================================

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None


try:
    from docx import Document
except ImportError:
    Document = None


try:
    from PIL import Image, ImageEnhance, ImageFilter
except ImportError:
    Image = None
    ImageEnhance = None
    ImageFilter = None


try:
    import pytesseract
except ImportError:
    pytesseract = None


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

MAX_FILE_MB = 10
MAX_FILE_BYTES = (
    MAX_FILE_MB * 1024 * 1024
)

ALLOWED_EXTENSIONS = {
    "pdf",
    "doc",
    "docx",
    "jpg",
    "jpeg",
    "png",
}


# ============================================================
# FASTAPI
# ============================================================

app = FastAPI(
    title="ResumeAI API",
    description="AI-powered Resume Screening API",
    version="4.0.0",
)


# ============================================================
# CORS
# ============================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# MEMORY STORAGE
# ============================================================

screenings: Dict[str, Dict[str, Any]] = {}

lock = threading.Lock()


# ============================================================
# SKILL DATABASE
# ============================================================

SKILL_DATABASE = {

    # Programming
    "python",
    "java",
    "javascript",
    "typescript",
    "c",
    "c++",
    "c#",
    "go",
    "golang",
    "rust",
    "kotlin",
    "swift",
    "dart",
    "php",
    "ruby",

    # Frontend
    "html",
    "css",
    "react",
    "react.js",
    "angular",
    "vue",
    "vue.js",
    "next.js",
    "nextjs",
    "node.js",
    "nodejs",
    "express",
    "express.js",

    # Backend
    "flask",
    "fastapi",
    "django",
    "spring",
    "spring boot",
    "rest api",
    "restful api",
    "graphql",

    # Database
    "mysql",
    "postgresql",
    "postgres",
    "sqlite",
    "mongodb",
    "firebase",
    "redis",
    "oracle",
    "sql",
    "nosql",

    # Cloud
    "aws",
    "amazon web services",
    "azure",
    "google cloud",
    "gcp",
    "docker",
    "kubernetes",

    # DevOps
    "linux",
    "bash",
    "jenkins",
    "ci/cd",
    "git",
    "github",
    "gitlab",

    # AI / ML
    "artificial intelligence",
    "ai",
    "machine learning",
    "deep learning",
    "data science",
    "data analysis",
    "numpy",
    "pandas",
    "scikit-learn",
    "sklearn",
    "tensorflow",
    "pytorch",
    "keras",
    "opencv",
    "nlp",
    "natural language processing",
    "computer vision",
    "llm",
    "generative ai",
    "generative artificial intelligence",
    "transformers",
    "hugging face",
    "yolo",

    # Visualization
    "matplotlib",
    "seaborn",
    "power bi",
    "tableau",
    "excel",

    # Mobile
    "flutter",
    "android",
    "android development",
    "ios",
    "react native",

    # Testing
    "pytest",
    "unit testing",
    "selenium",
    "postman",

    # General
    "api",
    "agile",
    "scrum",
    "communication",
    "leadership",
    "problem solving",
    "teamwork",
}


# ============================================================
# SKILL ALIASES
# ============================================================

SKILL_ALIASES = {

    "python3": "python",
    "python 3": "python",
    "py": "python",

    "js": "javascript",
    "javascript es6": "javascript",

    "ts": "typescript",

    "reactjs": "react",
    "react js": "react",

    "angularjs": "angular",
    "angular js": "angular",

    "vuejs": "vue",
    "vue js": "vue",

    "nextjs": "next.js",
    "next js": "next.js",

    "node": "node.js",
    "nodejs": "node.js",
    "node js": "node.js",

    "expressjs": "express",
    "express js": "express",

    "fast api": "fastapi",
    "fast-api": "fastapi",

    "postgres": "postgresql",
    "postgres db": "postgresql",
    "postgres database": "postgresql",
    "postgre sql": "postgresql",

    "mongo": "mongodb",
    "mongo db": "mongodb",

    "google cloud platform": "google cloud",
    "gcp": "google cloud",

    "aws cloud": "aws",

    "ml": "machine learning",
    "ai/ml": "machine learning",

    "sklearn": "scikit-learn",
    "scikit learn": "scikit-learn",

    "cv": "computer vision",

    "nlp": "natural language processing",

    "powerbi": "power bi",
    "power-bi": "power bi",

    "problem-solving": "problem solving",
}


# ============================================================
# ROLE DATABASE
# ============================================================

ROLE_KEYWORDS = {

    "software engineer": [
        "software engineer",
        "software developer",
        "software development",
        "software engineering",
    ],

    "data scientist": [
        "data scientist",
        "data science",
        "predictive modeling",
    ],

    "machine learning engineer": [
        "machine learning engineer",
        "ml engineer",
        "machine learning",
        "deep learning",
        "model development",
    ],

    "ai engineer": [
        "ai engineer",
        "artificial intelligence engineer",
        "artificial intelligence",
        "generative ai",
        "llm",
    ],

    "frontend developer": [
        "frontend developer",
        "front end developer",
        "frontend engineer",
        "react developer",
        "ui developer",
    ],

    "backend developer": [
        "backend developer",
        "back end developer",
        "backend engineer",
        "api developer",
        "server-side developer",
    ],

    "full stack developer": [
        "full stack developer",
        "full-stack developer",
        "fullstack developer",
    ],

    "flutter developer": [
        "flutter developer",
        "flutter",
        "dart developer",
        "mobile developer",
    ],
}


# ============================================================
# PYDANTIC MODELS
# ============================================================

class ScreeningOptions(BaseModel):

    minMatchScore: float = 60

    skillSensitivity: str = "balanced"

    aiExplanation: bool = True


class CreateScreeningRequest(BaseModel):

    jobDescription: str

    options: Optional[
        ScreeningOptions
    ] = None


# ============================================================
# BASIC HELPERS
# ============================================================

def generate_id() -> str:

    return uuid.uuid4().hex[:12]


def get_extension(
    filename: str,
) -> str:

    if (
        not filename
        or "."
        not in filename
    ):
        return ""

    return filename.rsplit(
        ".",
        1
    )[1].lower()


def allowed_file(
    filename: str,
) -> bool:

    return (
        get_extension(filename)
        in ALLOWED_EXTENSIONS
    )


def safe_float(
    value,
    default=0.0,
) -> float:

    try:
        return float(value)

    except (
        TypeError,
        ValueError,
    ):
        return default


def clamp(
    value: float,
    minimum: float = 0,
    maximum: float = 100,
) -> float:

    return max(
        minimum,
        min(
            maximum,
            safe_float(value),
        ),
    )


def clean_text(
    text: str,
) -> str:

    if not text:
        return ""

    text = text.replace(
        "\x00",
        " ",
    )

    text = re.sub(
        r"\r\n?",
        "\n",
        text,
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def normalize_text(
    text: str,
) -> str:

    text = str(
        text or ""
    ).lower()

    text = text.replace(
        "–",
        "-",
    )

    text = text.replace(
        "—",
        "-",
    )

    text = text.replace(
        "•",
        " ",
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def canonical_skill(
    skill: str,
) -> str:

    skill = normalize_text(
        skill
    )

    if not skill:
        return ""

    if skill in SKILL_ALIASES:

        skill = SKILL_ALIASES[
            skill
        ]

    return normalize_text(
        skill
    )


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_text(
    file_path: Path,
) -> str:

    if PdfReader is None:

        print(
            "[ERROR] PyPDF2 is not installed."
        )

        return ""

    try:

        reader = PdfReader(
            str(file_path)
        )

        pages = []

        for page in reader.pages:

            try:

                text = (
                    page.extract_text()
                    or ""
                )

                if text.strip():

                    pages.append(text)

            except Exception as exc:

                print(
                    "[PDF PAGE ERROR]",
                    exc,
                )

        result = clean_text(
            "\n".join(pages)
        )

        print(
            f"[PDF TEXT] {len(result)} characters"
        )

        return result

    except Exception as exc:

        print(
            "[PDF ERROR]",
            exc,
        )

        return ""


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(
    file_path: Path,
) -> str:

    if Document is None:

        print(
            "[ERROR] python-docx is not installed."
        )

        return ""

    try:

        document = Document(
            str(file_path)
        )

        parts = []

        for paragraph in (
            document.paragraphs
        ):

            if paragraph.text:

                parts.append(
                    paragraph.text
                )

        for table in document.tables:

            for row in table.rows:

                values = []

                for cell in row.cells:

                    if cell.text:

                        values.append(
                            cell.text
                        )

                if values:

                    parts.append(
                        " | ".join(values)
                    )

        return clean_text(
            "\n".join(parts)
        )

    except Exception as exc:

        print(
            "[DOCX ERROR]",
            exc,
        )

        return ""


# ============================================================
# OLD DOC EXTRACTION
# ============================================================

def extract_doc_text(
    file_path: Path,
) -> str:

    try:

        result = subprocess.run(
            [
                "antiword",
                str(file_path),
            ],
            capture_output=True,
            text=True,
            timeout=20,
        )

        if (
            result.returncode == 0
            and result.stdout
        ):

            return clean_text(
                result.stdout
            )

    except Exception as exc:

        print(
            "[DOC ERROR]",
            exc,
        )

    print(
        "[WARNING] Could not extract .doc. "
        "Install antiword or convert to .docx."
    )

    return ""


# ============================================================
# IMAGE OCR
# ============================================================

def extract_image_text(
    file_path: Path,
) -> str:

    if Image is None:

        print(
            "[ERROR] Pillow is not installed."
        )

        return ""

    if pytesseract is None:

        print(
            "[ERROR] pytesseract is not installed."
        )

        return ""

    try:

        image = Image.open(
            file_path
        )

        if image.mode != "RGB":

            image = image.convert(
                "RGB"
            )

        try:

            image = image.resize(
                (
                    image.width * 2,
                    image.height * 2,
                )
            )

            if ImageEnhance:

                image = (
                    ImageEnhance.Contrast(
                        image
                    ).enhance(1.5)
                )

            if ImageFilter:

                image = image.filter(
                    ImageFilter.SHARPEN
                )

        except Exception:
            pass

        text = pytesseract.image_to_string(
            image,
            config="--psm 6",
        )

        return clean_text(
            text
        )

    except Exception as exc:

        print(
            "[OCR ERROR]",
            exc,
        )

        return ""


# ============================================================
# GENERAL RESUME EXTRACTION
# ============================================================

def extract_resume_text(
    file_path: Path,
) -> str:

    extension = get_extension(
        file_path.name
    )

    if extension == "pdf":

        return extract_pdf_text(
            file_path
        )

    if extension == "docx":

        return extract_docx_text(
            file_path
        )

    if extension == "doc":

        return extract_doc_text(
            file_path
        )

    if extension in {
        "jpg",
        "jpeg",
        "png",
    }:

        return extract_image_text(
            file_path
        )

    return ""


# ============================================================
# SKILL EXTRACTION
# ============================================================

def extract_skills(
    text: str,
) -> List[str]:

    normalized = normalize_text(
        text
    )

    if not normalized:

        return []

    found = {}

    # --------------------------------------------------------
    # Database skills
    # --------------------------------------------------------

    for skill in SKILL_DATABASE:

        canonical = canonical_skill(
            skill
        )

        if not canonical:
            continue

        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(canonical)
            + r"(?![a-z0-9])"
        )

        try:

            if re.search(
                pattern,
                normalized,
                re.IGNORECASE,
            ):

                found[
                    canonical
                ] = skill

        except re.error:

            if canonical in normalized:

                found[
                    canonical
                ] = skill

    # --------------------------------------------------------
    # Aliases
    # --------------------------------------------------------

    for alias, canonical in (
        SKILL_ALIASES.items()
    ):

        alias_normalized = (
            normalize_text(alias)
        )

        canonical_normalized = (
            canonical_skill(canonical)
        )

        if not alias_normalized:
            continue

        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(
                alias_normalized
            )
            + r"(?![a-z0-9])"
        )

        try:

            if re.search(
                pattern,
                normalized,
                re.IGNORECASE,
            ):

                display = canonical

                for db_skill in (
                    SKILL_DATABASE
                ):

                    if (
                        canonical_skill(
                            db_skill
                        )
                        == canonical_normalized
                    ):

                        display = db_skill
                        break

                found[
                    canonical_normalized
                ] = display

        except re.error:
            pass

    return sorted(
        set(found.values()),
        key=lambda x: x.lower(),
    )


# ============================================================
# NAME
# ============================================================

def extract_name(
    text: str,
    filename: str = "",
) -> str:

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    # Explicit name
    patterns = [
        r"^(?:name|candidate name|full name)"
        r"\s*[:\-]\s*(.+)$",

        r"^candidate"
        r"\s*[:\-]\s*(.+)$",
    ]

    for line in lines[:30]:

        for pattern in patterns:

            match = re.match(
                pattern,
                line,
                re.IGNORECASE,
            )

            if match:

                name = (
                    match.group(1)
                    .strip()
                )

                if (
                    2
                    <= len(
                        name.split()
                    )
                    <= 6
                ):

                    return name

    blocked = {
        "resume",
        "curriculum",
        "vitae",
        "cv",
        "profile",
        "summary",
        "objective",
        "experience",
        "education",
        "skills",
        "skill",
        "contact",
        "email",
        "phone",
        "mobile",
        "address",
        "location",
        "developer",
        "engineer",
        "student",
        "projects",
        "project",
        "internship",
        "intern",
        "linkedin",
        "github",
        "portfolio",
    }

    for line in lines[:20]:

        clean = re.sub(
            r"[^A-Za-z .'-]",
            "",
            line,
        ).strip()

        words = clean.split()

        if len(words) < 2:
            continue

        if len(words) > 5:
            continue

        lower = clean.lower()

        if any(
            word in lower
            for word in blocked
        ):
            continue

        valid = all(
            re.match(
                r"^[A-Za-z][A-Za-z.'-]*$",
                word,
            )
            for word in words
        )

        if valid:

            return clean.title()

    # Filename fallback
    if filename:

        stem = Path(
            filename
        ).stem

        stem = re.sub(
            r"^[a-f0-9]{6,16}_",
            "",
            stem,
            flags=re.IGNORECASE,
        )

        stem = re.sub(
            r"whatsapp image\s*",
            "",
            stem,
            flags=re.IGNORECASE,
        )

        stem = re.sub(
            r"\b(resume|cv|final|updated|new)\b",
            "",
            stem,
            flags=re.IGNORECASE,
        )

        stem = re.sub(
            r"[_\-]+",
            " ",
            stem,
        )

        stem = re.sub(
            r"\s+",
            " ",
            stem,
        ).strip()

        if (
            stem
            and not re.fullmatch(
                r"[\d .:_-]+",
                stem,
            )
        ):

            return stem.title()

    return "Unknown Candidate"


# ============================================================
# LOCATION
# ============================================================

def extract_location(
    text: str,
) -> str:

    patterns = [

        r"(?:location|address|city)"
        r"\s*[:\-]\s*([^\n]+)",

        r"(?:based in|located in)"
        r"\s+([A-Za-z ,.-]{3,80})",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE,
        )

        if match:

            return (
                re.sub(
                    r"\s{2,}",
                    " ",
                    match.group(1),
                )
                .strip()[:100]
            )

    cities = [
        "Dehradun",
        "Delhi",
        "New Delhi",
        "Noida",
        "Gurugram",
        "Gurgaon",
        "Haridwar",
        "Rishikesh",
        "Chandigarh",
        "Mumbai",
        "Pune",
        "Bangalore",
        "Bengaluru",
        "Hyderabad",
        "Chennai",
        "Kolkata",
        "Jaipur",
        "Lucknow",
        "Ahmedabad",
        "Indore",
        "Kanpur",
        "Agra",
        "Hathras",
    ]

    lower = text.lower()

    for city in cities:

        if city.lower() in lower:

            return city

    return "Not specified"


# ============================================================
# EXPERIENCE
# ============================================================

def extract_year_number(
    text: str,
) -> Optional[float]:

    if not text:

        return None

    normalized = normalize_text(
        text
    )

    patterns = [

        r"(\d+(?:\.\d+)?)\+?\s*"
        r"(?:years?|yrs?)"
        r"(?:\s+of)?\s+experience",

        r"experience\s*[:\-]?\s*"
        r"(\d+(?:\.\d+)?)\+?\s*"
        r"(?:years?|yrs?)",

        r"(\d+(?:\.\d+)?)\+?\s*"
        r"(?:years?|yrs?)"
        r"\s+(?:in|of)\s+",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            normalized,
            re.IGNORECASE,
        )

        if match:

            return safe_float(
                match.group(1)
            )

    return None


def extract_experience(
    text: str,
) -> str:

    years = extract_year_number(
        text
    )

    if years is not None:

        if years == int(years):

            return (
                f"{int(years)} years"
            )

        return (
            f"{years:g} years"
        )

    if re.search(
        r"\bintern(ship)?\b",
        text,
        re.IGNORECASE,
    ):

        return "Internship experience"

    if re.search(
        r"\bfresher\b|"
        r"\bfresh graduate\b|"
        r"\brecent graduate\b",
        text,
        re.IGNORECASE,
    ):

        return "Fresher"

    return "Not specified"


# ============================================================
# ROLE
# ============================================================

def extract_role(
    text: str,
) -> str:

    normalized = normalize_text(
        text
    )

    for role, keywords in (
        ROLE_KEYWORDS.items()
    ):

        for keyword in keywords:

            if normalize_text(
                keyword
            ) in normalized:

                return role.title()

    return "Candidate"


# ============================================================
# EMAIL
# ============================================================

def extract_email(
    text: str,
) -> str:

    match = re.search(
        r"\b[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\."
        r"[A-Za-z]{2,}\b",
        text,
    )

    if match:

        return match.group(0)

    return ""


# ============================================================
# PHONE
# ============================================================

def extract_phone(
    text: str,
) -> str:

    patterns = [

        r"(?<!\d)"
        r"(?:\+91[\s-]?)?"
        r"[6-9]\d{9}"
        r"(?!\d)",

        r"(?<!\d)"
        r"\+91[\s-]?[6-9]\d{9}"
        r"(?!\d)",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
        )

        if match:

            return match.group(0)

    return ""


# ============================================================
# JOB DESCRIPTION ANALYSIS
# ============================================================

def analyze_job_description(
    job_description: str,
) -> Dict[str, Any]:

    skills = extract_skills(
        job_description
    )

    return {
        "skills": skills,

        "role": extract_role(
            job_description
        ),

        "requiredYears":
            extract_year_number(
                job_description
            ),

        "skillCount":
            len(skills),
    }


# ============================================================
# CANDIDATE ANALYSIS
# ============================================================

def analyze_candidate(
    resume_text: str,
    filename: str,
) -> Dict[str, Any]:

    return {

        "name": extract_name(
            resume_text,
            filename,
        ),

        "role": extract_role(
            resume_text,
        ),

        "experience": extract_experience(
            resume_text,
        ),

        "location": extract_location(
            resume_text,
        ),

        "email": extract_email(
            resume_text,
        ),

        "phone": extract_phone(
            resume_text,
        ),

        "skills": extract_skills(
            resume_text,
        ),
    }


# ============================================================
# SKILL MATCHING
# ============================================================

def calculate_skill_match(
    jd_skills: List[str],
    candidate_skills: List[str],
    sensitivity: str = "balanced",
) -> Dict[str, Any]:

    jd_map = {}

    for skill in jd_skills or []:

        canonical = canonical_skill(
            skill
        )

        if canonical:

            jd_map[
                canonical
            ] = skill


    candidate_set = set()

    for skill in candidate_skills or []:

        canonical = canonical_skill(
            skill
        )

        if canonical:

            candidate_set.add(
                canonical
            )


    # --------------------------------------------------------
    # No detected JD skills
    # --------------------------------------------------------

    if not jd_map:

        return {
            "score": 50.0,
            "matched": [],
            "missing": [],
        }


    matched = set()


    # --------------------------------------------------------
    # Exact / related matching
    # --------------------------------------------------------

    for jd_skill in jd_map:

        if jd_skill in candidate_set:

            matched.add(
                jd_skill
            )

            continue


        for candidate_skill in (
            candidate_set
        ):

            if (
                jd_skill == candidate_skill
                or
                jd_skill in candidate_skill
                or
                candidate_skill in jd_skill
            ):

                matched.add(
                    jd_skill
                )

                break


    missing = (
        set(jd_map.keys())
        - matched
    )


    matched_display = [
        jd_map[x]
        for x in matched
    ]

    missing_display = [
        jd_map[x]
        for x in missing
    ]


    score = (
        len(matched)
        /
        len(jd_map)
    ) * 100


    sensitivity = str(
        sensitivity or "balanced"
    ).lower()


    if sensitivity == "strict":

        score *= 0.95

    elif sensitivity == "flexible":

        score = min(
            100,
            score * 1.05 + 3,
        )


    return {

        "score": round(
            clamp(score),
            1,
        ),

        "matched": sorted(
            matched_display,
            key=lambda x: x.lower(),
        ),

        "missing": sorted(
            missing_display,
            key=lambda x: x.lower(),
        ),
    }


# ============================================================
# KEYWORD MATCH
# ============================================================

STOP_WORDS = {
    "the",
    "and",
    "for",
    "with",
    "that",
    "this",
    "from",
    "into",
    "your",
    "you",
    "are",
    "will",
    "have",
    "has",
    "our",
    "their",
    "they",
    "them",
    "job",
    "role",
    "work",
    "working",
    "candidate",
    "required",
    "requirements",
    "preferred",
    "using",
    "looking",
    "years",
    "year",
    "experience",
    "responsibilities",
    "skills",
    "should",
    "must",
    "ability",
    "knowledge",
    "strong",
    "good",
    "team",
    "development",
    "developer",
}


def calculate_keyword_match(
    job_description: str,
    resume_text: str,
) -> float:

    jd = normalize_text(
        job_description
    )

    resume = normalize_text(
        resume_text
    )

    if not jd or not resume:

        return 0.0


    jd_words = set(
        word
        for word in re.findall(
            r"[a-zA-Z0-9+#.-]+",
            jd,
        )
        if len(word) >= 4
        and word not in STOP_WORDS
    )


    resume_words = set(
        re.findall(
            r"[a-zA-Z0-9+#.-]+",
            resume,
        )
    )


    if not jd_words:

        return 50.0


    matched = (
        jd_words
        &
        resume_words
    )


    return round(
        clamp(
            (
                len(matched)
                /
                len(jd_words)
            )
            * 100
        ),
        1,
    )


# ============================================================
# PHRASE MATCH
# ============================================================

def calculate_phrase_match(
    job_description: str,
    resume_text: str,
) -> float:

    jd = normalize_text(
        job_description
    )

    resume = normalize_text(
        resume_text
    )

    if not jd or not resume:

        return 0.0


    phrases = [
        "machine learning",
        "deep learning",
        "artificial intelligence",
        "data science",
        "data analysis",
        "web development",
        "mobile development",
        "backend development",
        "frontend development",
        "full stack",
        "software development",
        "rest api",
        "cloud computing",
        "database management",
        "problem solving",
        "object oriented programming",
        "computer vision",
        "natural language processing",
        "generative ai",
    ]


    jd_phrases = [
        phrase
        for phrase in phrases
        if phrase in jd
    ]


    if not jd_phrases:

        return 50.0


    matched = [
        phrase
        for phrase in jd_phrases
        if phrase in resume
    ]


    return round(
        (
            len(matched)
            /
            len(jd_phrases)
        )
        * 100,
        1,
    )


# ============================================================
# EXPERIENCE SCORE
# ============================================================

def calculate_experience_score(
    job_description: str,
    resume_text: str,
) -> float:

    required = extract_year_number(
        job_description
    )

    candidate = extract_year_number(
        resume_text
    )


    if required is None:

        return 100.0


    if required <= 0:

        return 100.0


    if candidate is None:

        return 50.0


    if candidate >= required:

        return 100.0


    return round(
        clamp(
            (
                candidate
                /
                required
            )
            * 100
        ),
        1,
    )


# ============================================================
# FINAL SCORE
# ============================================================

def calculate_final_score(
    skill_score: float,
    keyword_score: float,
    phrase_score: float,
    experience_score: float,
    job_description: str,
) -> float:

    skill_score = clamp(
        skill_score
    )

    keyword_score = clamp(
        keyword_score
    )

    phrase_score = clamp(
        phrase_score
    )

    experience_score = clamp(
        experience_score
    )


    required = extract_year_number(
        job_description
    )


    # --------------------------------------------------------
    # NO EXPERIENCE REQUIREMENT
    #
    # Skills       = 60%
    # Keywords     = 25%
    # Phrases      = 15%
    # --------------------------------------------------------

    if required is None:

        final_score = (
            skill_score * 0.60
            +
            keyword_score * 0.25
            +
            phrase_score * 0.15
        )


    # --------------------------------------------------------
    # EXPERIENCE REQUIREMENT
    #
    # Skills       = 55%
    # Keywords     = 20%
    # Phrases      = 10%
    # Experience   = 15%
    # --------------------------------------------------------

    else:

        final_score = (
            skill_score * 0.55
            +
            keyword_score * 0.20
            +
            phrase_score * 0.10
            +
            experience_score * 0.15
        )


    return round(
        clamp(final_score),
        1,
    )


# ============================================================
# MATCH ANALYSIS
# ============================================================

def calculate_candidate_match(
    job_description: str,
    resume_text: str,
    jd_skills: List[str],
    candidate_skills: List[str],
    sensitivity: str = "balanced",
) -> Dict[str, Any]:

    skill_result = calculate_skill_match(
        jd_skills,
        candidate_skills,
        sensitivity,
    )


    skill_score = safe_float(
        skill_result.get(
            "score",
            0,
        )
    )


    keyword_score = (
        calculate_keyword_match(
            job_description,
            resume_text,
        )
    )


    phrase_score = (
        calculate_phrase_match(
            job_description,
            resume_text,
        )
    )


    experience_score = (
        calculate_experience_score(
            job_description,
            resume_text,
        )
    )


    final_score = (
        calculate_final_score(
            skill_score,
            keyword_score,
            phrase_score,
            experience_score,
            job_description,
        )
    )


    return {

        "score": final_score,

        "matchScore": final_score,

        "match_score": final_score,

        "skillScore": skill_score,

        "keywordScore": keyword_score,

        "phraseScore": phrase_score,

        "experienceScore":
            experience_score,

        "matchedSkills":
            skill_result[
                "matched"
            ],

        "missingSkills":
            skill_result[
                "missing"
            ],
    }


# ============================================================
# EXPLANATION
# ============================================================

def generate_explanation(
    candidate: Dict[str, Any],
    match_result: Dict[str, Any],
) -> str:

    score = safe_float(
        match_result.get(
            "score",
            0,
        )
    )


    matched = match_result.get(
        "matchedSkills",
        [],
    )


    missing = match_result.get(
        "missingSkills",
        [],
    )


    if score >= 80:

        level = "Strong match"

    elif score >= 60:

        level = "Good match"

    elif score >= 40:

        level = "Partial match"

    else:

        level = "Low match"


    parts = [

        f"{level}.",

        (
            f"The candidate matches "
            f"{len(matched)} identified "
            f"job skills."
        ),
    ]


    if matched:

        parts.append(
            "Matching skills: "
            +
            ", ".join(
                matched[:8]
            )
            +
            "."
        )


    if missing:

        parts.append(
            "Missing skills: "
            +
            ", ".join(
                missing[:8]
            )
            +
            "."
        )


    parts.append(
        "Experience: "
        +
        str(
            candidate.get(
                "experience",
                "Not specified",
            )
        )
        +
        "."
    )


    return " ".join(parts)


# ============================================================
# PROCESS ONE RESUME
# ============================================================

def process_candidate(
    file_path: Path,
    job_analysis: Dict[str, Any],
    options: Dict[str, Any],
) -> Dict[str, Any]:

    filename = file_path.name


    print("=" * 60)

    print(
        f"[PROCESSING] {filename}"
    )


    # --------------------------------------------------------
    # Extract text
    # --------------------------------------------------------

    resume_text = extract_resume_text(
        file_path
    )


    print(
        f"[TEXT LENGTH] "
        f"{len(resume_text)} characters"
    )


    if not resume_text:

        print(
            "[WARNING] No text extracted."
        )


    # --------------------------------------------------------
    # Candidate information
    # --------------------------------------------------------

    candidate = analyze_candidate(
        resume_text,
        filename,
    )


    print(
        "[CANDIDATE]",
        candidate["name"],
    )

    print(
        "[CANDIDATE SKILLS]",
        candidate["skills"],
    )


    # --------------------------------------------------------
    # IMPORTANT: calculate ALL score components
    # --------------------------------------------------------

    match_result = (
        calculate_candidate_match(
            job_description=
                options.get(
                    "jobDescription",
                    "",
                ),

            resume_text=
                resume_text,

            jd_skills=
                job_analysis.get(
                    "skills",
                    [],
                ),

            candidate_skills=
                candidate.get(
                    "skills",
                    [],
                ),

            sensitivity=
                options.get(
                    "skillSensitivity",
                    "balanced",
                ),
        )
    )


    final_score = safe_float(
        match_result.get(
            "score",
            0,
        )
    )


    explanation = (
        generate_explanation(
            candidate,
            match_result,
        )
    )


    # --------------------------------------------------------
    # Debug
    # --------------------------------------------------------

    print(
        "[JD SKILLS]",
        job_analysis.get(
            "skills",
            [],
        ),
    )

    print(
        "[MATCHED]",
        match_result.get(
            "matchedSkills",
            [],
        ),
    )

    print(
        "[MISSING]",
        match_result.get(
            "missingSkills",
            [],
        ),
    )

    print(
        "[SKILL SCORE]",
        match_result.get(
            "skillScore",
            0,
        ),
    )

    print(
        "[KEYWORD SCORE]",
        match_result.get(
            "keywordScore",
            0,
        ),
    )

    print(
        "[PHRASE SCORE]",
        match_result.get(
            "phraseScore",
            0,
        ),
    )

    print(
        "[EXPERIENCE SCORE]",
        match_result.get(
            "experienceScore",
            0,
        ),
    )

    print(
        "[FINAL SCORE]",
        final_score,
    )

    print("=" * 60)


    return {

        "name":
            candidate["name"],

        "role":
            candidate["role"],

        "experience":
            candidate["experience"],

        "location":
            candidate["location"],

        "email":
            candidate["email"],

        "phone":
            candidate["phone"],

        "skills":
            candidate["skills"],

        "matchedSkills":
            match_result[
                "matchedSkills"
            ],

        "missingSkills":
            match_result[
                "missingSkills"
            ],

        "skillScore":
            match_result[
                "skillScore"
            ],

        "keywordScore":
            match_result[
                "keywordScore"
            ],

        "phraseScore":
            match_result[
                "phraseScore"
            ],

        "experienceScore":
            match_result[
                "experienceScore"
            ],

        # Frontend compatibility
        "matchScore":
            final_score,

        "score":
            final_score,

        "match_score":
            final_score,

        "explanation":
            explanation,

        "resumeFilename":
            filename,

        "resumeUrl":
            None,

        # Internal only
        "resumeText":
            resume_text,
    }


# ============================================================
# SCREENING UPDATE
# ============================================================

def update_screening(
    screening_id: str,
    **values,
):

    with lock:

        screening = screenings.get(
            screening_id
        )

        if not screening:
            return

        screening.update(
            values
        )


# ============================================================
# BACKGROUND SCREENING
# ============================================================

def run_screening(
    screening_id: str,
):

    try:

        with lock:

            screening = screenings.get(
                screening_id
            )

            if not screening:
                return

            job_description = (
                screening[
                    "jobDescription"
                ]
            )

            options = dict(
                screening[
                    "options"
                ]
            )

            files = list(
                screening[
                    "files"
                ]
            )


        # ----------------------------------------------------
        # STEP 1
        # ----------------------------------------------------

        update_screening(
            screening_id,

            currentStep=
                "readingJobDescription",

            progress=10,

            statusMessage=
                "Reading job description...",
        )


        job_analysis = (
            analyze_job_description(
                job_description
            )
        )


        options[
            "jobDescription"
        ] = job_description


        print("=" * 60)

        print(
            "[JOB ROLE]",
            job_analysis[
                "role"
            ],
        )

        print(
            "[JOB SKILLS]",
            job_analysis[
                "skills"
            ],
        )

        print(
            "[REQUIRED YEARS]",
            job_analysis[
                "requiredYears"
            ],
        )

        print("=" * 60)


        # ----------------------------------------------------
        # STEP 2
        # ----------------------------------------------------

        update_screening(
            screening_id,

            currentStep=
                "uploadingCvs",

            progress=25,

            statusMessage=
                "Preparing resumes...",
        )


        # ----------------------------------------------------
        # STEP 3
        # ----------------------------------------------------

        candidates = []

        total_files = len(
            files
        )


        for index, file_info in (
            enumerate(files)
        ):

            filename = file_info[
                "filename"
            ]

            path = Path(
                file_info["path"]
            )


            progress = (
                30
                +
                (
                    (
                        index + 1
                    )
                    /
                    max(
                        total_files,
                        1,
                    )
                )
                * 30
            )


            update_screening(
                screening_id,

                currentStep=
                    "extractingInformation",

                progress=progress,

                statusMessage=
                    f"Reading {filename}...",
            )


            result = process_candidate(
                path,
                job_analysis,
                options,
            )


            result["resumeUrl"] = (
                "/api/files/"
                +
                screening_id
                +
                "/"
                +
                filename
            )


            candidates.append(
                result
            )


        # ----------------------------------------------------
        # STEP 4
        # ----------------------------------------------------

        update_screening(
            screening_id,

            currentStep=
                "matchingSkills",

            progress=70,

            statusMessage=
                "Matching candidate skills...",
        )


        # ----------------------------------------------------
        # STEP 5
        # ----------------------------------------------------

        update_screening(
            screening_id,

            currentStep=
                "calculatingScore",

            progress=85,

            statusMessage=
                "Calculating candidate scores...",
        )


        candidates.sort(
            key=lambda candidate:
                safe_float(
                    candidate.get(
                        "matchScore",
                        candidate.get(
                            "score",
                            0,
                        ),
                    )
                ),
            reverse=True,
        )


        # ----------------------------------------------------
        # COMPLETE
        # ----------------------------------------------------

        update_screening(
            screening_id,

            candidates=candidates,

            status="completed",

            progress=100,

            currentStep="completed",

            statusMessage=
                "Screening completed successfully.",

            completedAt=
                datetime.utcnow().isoformat(),
        )


        print("=" * 60)

        print(
            "[SUCCESS] Screening completed:",
            screening_id,
        )

        print(
            "[CANDIDATES]",
            len(candidates),
        )

        print("=" * 60)


    except Exception as exc:

        traceback.print_exc()

        update_screening(
            screening_id,

            status="error",

            currentStep="error",

            statusMessage=
                f"Processing failed: {exc}",

            error=str(exc),
        )


# ============================================================
# ROOT
# ============================================================

@app.get("/")
def root():

    return {

        "service":
            "ResumeAI Backend",

        "status":
            "running",

        "framework":
            "FastAPI",

        "version":
            "4.0.0",

        "docs":
            "/docs",

        "health":
            "/api/health",
    }


# ============================================================
# HEALTH
# ============================================================
@app.get("/api/health")
def health():

    return {
        "status": "ok",

        "service": "ResumeAI Backend",

        "framework": "FastAPI",

        "version": "4.0.0",

        "pdfExtractor":
            "PyPDF2"
            if PdfReader
            else None,

        "docxExtractor":
            bool(Document),

        "ocr":
            bool(pytesseract),

        # "pymupdf":
            # False,
    }

# ============================================================
# CREATE SCREENING
# ============================================================

@app.post(
    "/api/screenings"
)
def create_screening(
    data: CreateScreeningRequest,
):

    job_description = (
        data.jobDescription
        or ""
    ).strip()


    if not job_description:

        raise HTTPException(
            status_code=400,
            detail=
                "Job description is required.",
        )


    options = (
        data.options
        or ScreeningOptions()
    )


    screening_id = generate_id()


    screening = {

        "id":
            screening_id,

        "createdAt":
            datetime.utcnow().isoformat(),

        "jobDescription":
            job_description,

        "options": {

            "minMatchScore":
                options.minMatchScore,

            "skillSensitivity":
                options.skillSensitivity,

            "aiExplanation":
                options.aiExplanation,
        },

        "status":
            "created",

        "progress":
            0,

        "currentStep":
            None,

        "statusMessage":
            "Screening created.",

        "error":
            None,

        "files":
            [],

        "candidates":
            [],

        "completedAt":
            None,
    }


    with lock:

        screenings[
            screening_id
        ] = screening


    print(
        "[SCREENING CREATED]",
        screening_id,
    )


    return {

        "screeningId":
            screening_id,

        "id":
            screening_id,
    }


# ============================================================
# UPLOAD RESUMES
# ============================================================

@app.post(
    "/api/screenings/{screening_id}/resumes"
)
async def upload_resumes(
    screening_id: str,
    background_tasks:
        BackgroundTasks,

    resumes:
        List[UploadFile] = File(...),
):

    with lock:

        screening = screenings.get(
            screening_id
        )


    if not screening:

        raise HTTPException(
            status_code=404,
            detail=
                "Screening not found.",
        )


    if not resumes:

        raise HTTPException(
            status_code=400,
            detail=
                "No resume files were uploaded.",
        )


    screening_dir = (
        UPLOAD_DIR
        /
        screening_id
    )


    screening_dir.mkdir(
        parents=True,
        exist_ok=True,
    )


    saved_files = []


    for uploaded_file in resumes:

        if not uploaded_file.filename:

            continue


        original_name = Path(
            uploaded_file.filename
        ).name


        if not allowed_file(
            original_name
        ):

            print(
                "[SKIPPED] Unsupported:",
                original_name,
            )

            continue


        safe_name = re.sub(
            r"[^A-Za-z0-9._-]",
            "_",
            original_name,
        )


        stored_name = (
            uuid.uuid4().hex[:8]
            +
            "_"
            +
            safe_name
        )


        save_path = (
            screening_dir
            /
            stored_name
        )


        try:

            content = await (
                uploaded_file.read()
            )


            if len(content) > (
                MAX_FILE_BYTES
            ):

                print(
                    "[SKIPPED] Too large:",
                    original_name,
                )

                continue


            if len(content) == 0:

                print(
                    "[SKIPPED] Empty:",
                    original_name,
                )

                continue


            with open(
                save_path,
                "wb",
            ) as output_file:

                output_file.write(
                    content
                )


            saved_files.append({

                "filename":
                    stored_name,

                "originalFilename":
                    original_name,

                "path":
                    str(save_path),
            })


            print(
                "[UPLOADED]",
                original_name,
            )


        except Exception as exc:

            print(
                "[UPLOAD ERROR]",
                original_name,
                exc,
            )


    if not saved_files:

        raise HTTPException(
            status_code=400,

            detail={
                "message":
                    "No valid resume files were uploaded.",

                "allowedExtensions":
                    sorted(
                        ALLOWED_EXTENSIONS
                    ),

                "maxFileSizeMB":
                    MAX_FILE_MB,
            },
        )


    update_screening(
        screening_id,

        files=saved_files,

        status="processing",

        progress=5,

        currentStep=
            "uploadingCvs",

        statusMessage=
            "Files uploaded. "
            "Starting screening...",
    )


    background_tasks.add_task(
        run_screening,
        screening_id,
    )


    return {

        "message":
            "Resumes uploaded successfully.",

        "screeningId":
            screening_id,

        "fileCount":
            len(saved_files),
    }


# ============================================================
# STATUS
# ============================================================

@app.get(
    "/api/screenings/{screening_id}/status"
)
def screening_status(
    screening_id: str,
):

    with lock:

        screening = screenings.get(
            screening_id
        )


    if not screening:

        raise HTTPException(
            status_code=404,
            detail=
                "Screening not found.",
        )


    return {

        "id":
            screening["id"],

        "status":
            screening["status"],

        "progress":
            screening["progress"],

        "currentStep":
            screening["currentStep"],

        "statusMessage":
            screening["statusMessage"],

        "error":
            screening["error"],

        "total":
            len(
                screening["files"]
            ),

        "completed":
            len(
                screening["candidates"]
            ),

        "createdAt":
            screening["createdAt"],

        "completedAt":
            screening["completedAt"],
    }


# ============================================================
# GET CANDIDATE
# ============================================================

@app.get(
    "/api/screenings/{screening_id}/candidates/{index}"
)
def get_candidate(
    screening_id: str,
    index: int,
):

    with lock:

        screening = screenings.get(
            screening_id
        )

        if not screening:

            raise HTTPException(
                status_code=404,
                detail=
                    "Screening not found.",
            )

        candidates = list(
            screening["candidates"]
        )


    if (
        index < 0
        or index >= len(candidates)
    ):

        raise HTTPException(
            status_code=404,
            detail=
                "Candidate not found.",
        )


    candidate = dict(
        candidates[index]
    )


    candidate.pop(
        "resumeText",
        None,
    )


    score = safe_float(
        candidate.get(
            "matchScore",
            candidate.get(
                "score",
                0,
            ),
        )
    )


    candidate["matchScore"] = round(
        clamp(score),
        1,
    )


    candidate["score"] = (
        candidate["matchScore"]
    )


    candidate["match_score"] = (
        candidate["matchScore"]
    )


    candidate["index"] = index

    candidate["total"] = len(
        candidates
    )


    return candidate


# ============================================================
# SUMMARY
# ============================================================

@app.get(
    "/api/screenings/{screening_id}/summary"
)
def get_summary(
    screening_id: str,
):

    with lock:

        screening = screenings.get(
            screening_id
        )

        if not screening:

            raise HTTPException(
                status_code=404,
                detail=
                    "Screening not found.",
            )

        candidates = list(
            screening["candidates"]
        )

        options = dict(
            screening["options"]
        )


    normalized = []


    for candidate in candidates:

        item = dict(
            candidate
        )


        item.pop(
            "resumeText",
            None,
        )


        score = safe_float(
            item.get(
                "matchScore",
                item.get(
                    "score",
                    0,
                ),
            )
        )


        item["matchScore"] = round(
            clamp(score),
            1,
        )


        item["score"] = (
            item["matchScore"]
        )


        item["match_score"] = (
            item["matchScore"]
        )


        normalized.append(
            item
        )


    normalized.sort(
        key=lambda x:
            safe_float(
                x.get(
                    "matchScore",
                    0,
                )
            ),
        reverse=True,
    )


    total = len(
        normalized
    )


    minimum = safe_float(
        options.get(
            "minMatchScore",
            60,
        ),
        60,
    )


    strong_matches = [
        x
        for x in normalized
        if safe_float(
            x.get(
                "matchScore",
                0,
            )
        ) >= minimum
    ]


    scores = [
        safe_float(
            x.get(
                "matchScore",
                0,
            )
        )
        for x in normalized
    ]


    average = (
        sum(scores)
        /
        len(scores)
        if scores
        else 0
    )


    top = (
        normalized[0]
        if normalized
        else None
    )


    return {

        "screeningId":
            screening_id,

        "status":
            screening["status"],

        "total":
            total,

        "totalCandidates":
            total,

        "strongMatches":
            len(strong_matches),

        "averageMatch":
            round(
                average,
                1,
            ),

        "topMatch":
            round(
                safe_float(
                    top.get(
                        "matchScore",
                        0,
                    )
                )
                if top
                else 0,
                1,
            ),

        "topCandidate":
            {
                "name":
                    top.get(
                        "name",
                        "Unknown Candidate",
                    ),

                "score":
                    top.get(
                        "matchScore",
                        0,
                    ),

                "matchScore":
                    top.get(
                        "matchScore",
                        0,
                    ),

                "role":
                    top.get(
                        "role",
                        "Candidate",
                    ),
            }
            if top
            else None,

        "candidates":
            normalized,
    }


# ============================================================
# SERVE RESUME FILE
# ============================================================

@app.get(
    "/api/files/{screening_id}/{filename:path}"
)
def serve_file(
    screening_id: str,
    filename: str,
):

    screening_dir = (
        UPLOAD_DIR
        /
        screening_id
    )


    if not screening_dir.exists():

        raise HTTPException(
            status_code=404,
            detail=
                "Screening files not found.",
        )


    safe_filename = Path(
        filename
    ).name


    requested_file = (
        screening_dir
        /
        safe_filename
    )


    if not requested_file.exists():

        raise HTTPException(
            status_code=404,
            detail=
                "Resume file not found.",
        )


    return FileResponse(
        path=str(
            requested_file
        ),
        filename=safe_filename,
        content_disposition_type=
            "inline",
    )


# ============================================================
# LIST SCREENINGS
# ============================================================

@app.get(
    "/api/screenings"
)
def list_screenings():

    with lock:

        result = []

        for screening in (
            screenings.values()
        ):

            result.append({

                "id":
                    screening["id"],

                "createdAt":
                    screening["createdAt"],

                "status":
                    screening["status"],

                "progress":
                    screening["progress"],

                "total":
                    len(
                        screening["files"]
                    ),

                "completed":
                    len(
                        screening["candidates"]
                    ),
            })


    result.sort(
        key=lambda x:
            x["createdAt"],
        reverse=True,
    )


    return {
        "screenings":
            result
    }


# ============================================================
# DELETE SCREENING
# ============================================================

@app.delete(
    "/api/screenings/{screening_id}"
)
def delete_screening(
    screening_id: str,
):

    with lock:

        if screening_id not in screenings:

            raise HTTPException(
                status_code=404,
                detail=
                    "Screening not found.",
            )

        del screenings[
            screening_id
        ]


    screening_dir = (
        UPLOAD_DIR
        /
        screening_id
    )


    if screening_dir.exists():

        for file in (
            screening_dir.iterdir()
        ):

            try:

                if file.is_file():

                    file.unlink()

            except Exception:
                pass


        try:

            screening_dir.rmdir()

        except Exception:
            pass


    return {

        "message":
            "Screening deleted successfully.",

        "screeningId":
            screening_id,
    }


# ============================================================
# GLOBAL ERROR HANDLER
# ============================================================

@app.exception_handler(Exception)
async def global_exception_handler(
    request: Request,
    exc: Exception,
):

    traceback.print_exc()


    return JSONResponse(
        status_code=500,

        content={

            "message":
                "Internal server error.",

            "error":
                str(exc),
        },
    )


# ============================================================
# START SERVER
# ============================================================

if __name__ == "__main__":

    import uvicorn


    print("=" * 60)

    print(
        "ResumeAI FastAPI Backend"
    )

    print("=" * 60)

    print(
        f"Upload directory: "
        f"{UPLOAD_DIR}"
    )

    print(
        "Allowed files:",
        ", ".join(
            sorted(
                ALLOWED_EXTENSIONS
            )
        ),
    )

    print(
        "PDF extractor:",
        "PyPDF2"
        if PdfReader
        else "NOT INSTALLED",
    )

    print(
        "Server:",
        "http://127.0.0.1:5000",
    )

    print(
        "Swagger:",
        "http://127.0.0.1:5000/docs",
    )

    print(
        "Health:",
        "http://127.0.0.1:5000/api/health",
    )

    print("=" * 60)


    uvicorn.run(
        app,
        host="127.0.0.1",
        port=5000,
        reload=False,
    )